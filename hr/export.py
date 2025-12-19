import re
import openpyxl
from openpyxl.styles import Alignment, Font
from django.http import HttpResponse
from users.models import Resume, WorkExperience
from .educational_level import *
from .logs import *


# -------------------------------------------------
# Excel safety helper (CRITICAL)
# -------------------------------------------------
ILLEGAL_CHARACTERS_RE = re.compile(r'[\x00-\x08\x0B\x0C\x0E-\x1F]')


def clean_excel_value(value):
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub('', value)
    return value


def export_applications_to_excel(request, applications):
    """
    Export application data to an Excel file (Excel-safe).
    """
    create_log(request.user, "Exported Applications")

    data = []

    for application in applications:
        user = application.applicant

        try:
            resume = user.resume
        except Resume.DoesNotExist:
            resume = Resume.objects.create(
                user=user,
                full_name=f"{user.first_name} {user.last_name}",
                email_address=user.email,
            )

        basic_education = user.basic_education.all()[:2]
        further_studies = user.further_studies.all()[:3]
        work_experience = user.work_experiences.all()[:3]
        certifications = user.certifications.all()[:3]
        memberships = user.memberships.all()[:3]
        referees = user.referees.all()[:3]

        contacts = f"{resume.phone}\n{resume.email_address}"

        educational_level = resume.educational_level or get_user_highest_educational_level(
            user.id)
        educational_level_name = educational_level.name if educational_level else ''

        application_data = {
            'Username/Staff No.': user.username,
            'Full Name': resume.full_name,
            'Contact Details': contacts,
            'Gender': resume.gender or '',
            'Disability': 'Yes' if resume.disability else 'No',
            'Ethnicity': resume.ethnicity.name if resume.ethnicity else '',
            'Highest Educational Level': educational_level_name,
            'High School': '',
            'College/University': '',
            'Professional Certifications': '',
            'Professional Membership': '',
            'Work Experience': '',
            'Referees': '',
            'Date Applied': application.application_date.strftime("%Y-%m-%d %H:%M:%S"),
            'Reference Number': application.reference_number,
            'Status': 'Qualified' if application.qualify else 'Not Qualified',
            'Shortlisted': 'Yes' if application.shortlisted else 'No',
        }

        # ---------------- BASIC EDUCATION ----------------
        if basic_education:
            block = []
            for edu in basic_education:
                block.append(
                    f"School: {edu.name_of_the_school}\n"
                    f"Start Year: {edu.date_started}\n"
                    f"End Year: {edu.date_ended}\n"
                    f"Grade Attained: {edu.grade_attained}"
                )
            application_data['High School'] = "\n\n".join(block)

        # ---------------- FURTHER STUDIES ----------------
        if further_studies:
            block = []
            for study in further_studies:
                block.append(
                    f"Institution: {study.institution_name}\n"
                    f"Certification: {study.certifications.name if study.certifications else ''}\n"
                    f"Course: {study.course_undertaken}\n"
                    f"Start Date: {study.date_started.strftime('%Y-%m-%d') if study.date_started else ''}\n"
                    f"End Date: {study.date_ended.strftime('%Y-%m-%d') if study.date_ended else ''}\n"
                    f"Grade: {study.grade}"
                )
            application_data['College/University'] = "\n\n".join(block)

        # ---------------- CERTIFICATIONS ----------------
        if certifications:
            block = []
            for cert in certifications:
                block.append(
                    f"Name: {cert.name}\n"
                    f"Body: {cert.certifying_body.name if cert.certifying_body else ''}\n"
                    f"Date Attained: {cert.date_attained.strftime('%Y-%m-%d') if cert.date_attained else ''}"
                )
            application_data['Professional Certifications'] = "\n\n".join(
                block)

        # ---------------- MEMBERSHIPS ----------------
        if memberships:
            block = []
            for m in memberships:
                block.append(
                    f"Title: {m.membership_title}\n"
                    f"Number: {m.membership_number}\n"
                    f"Body: {m.membership_body}\n"
                    f"Date Joined: {m.date_joined.strftime('%Y-%m-%d') if m.date_joined else ''}"
                )
            application_data['Professional Membership'] = "\n\n".join(block)

        # ---------------- WORK EXPERIENCE ----------------
        if work_experience:
            block = []
            total_months = 0

            for exp in work_experience:
                start = exp.date_started.strftime(
                    '%Y-%m-%d') if exp.date_started else ''
                end = exp.date_ended.strftime(
                    '%Y-%m-%d') if exp.date_ended else 'In Progress'

                block.append(
                    f"Company: {exp.company_name}\n"
                    f"Position: {exp.position}\n"
                    f"Start Date: {start}\n"
                    f"End Date: {end}\n"
                    f"Address: {exp.company_address}\n"
                    f"Phone: {exp.company_phone}\n"
                    f"Responsibilities: {exp.responsibilities}"
                )

                if exp.date_started and exp.date_ended:
                    delta = exp.date_ended - exp.date_started
                    total_months += delta.days // 30

            years, months = divmod(total_months, 12)
            block.append(f"Total Experience: {years} years {months} months")

            application_data['Work Experience'] = "\n\n".join(block)

        # ---------------- REFEREES ----------------
        if referees:
            block = []
            for r in referees:
                block.append(
                    f"Name: {r.full_name}\n"
                    f"Organization: {r.organization}\n"
                    f"Designation: {r.designation}\n"
                    f"Phone: {r.phone}\n"
                    f"Email: {r.email}"
                )
            application_data['Referees'] = "\n\n".join(block)

        data.append(application_data)

    headers = [
        'Username/Staff No.', 'Full Name', 'Contact Details', 'Gender', 'Disability',
        'Ethnicity', 'Highest Educational Level', 'High School', 'College/University',
        'Professional Certifications', 'Professional Membership', 'Work Experience',
        'Referees', 'Date Applied', 'Reference Number', 'Status', 'Shortlisted'
    ]

    wb = openpyxl.Workbook()
    ws = wb.active

    title = clean_excel_value(
        f"{application.vacancy.title} / {application.vacancy.ref} Applications"
    )

    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = Font(size=14, bold=True)
    title_cell.alignment = Alignment(horizontal='center')

    for col, header in enumerate(headers, 1):
        ws.cell(row=2, column=col, value=header)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 28

    for row, record in enumerate(data, 3):
        for col, header in enumerate(headers, 1):
            value = clean_excel_value(record.get(header, ''))
            cell = ws.cell(row=row, column=col, value=value)
            cell.alignment = Alignment(wrap_text=True, vertical='top')

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = (
        f'attachment; filename="applications_for_'
        f'{application.vacancy.title}_{application.vacancy.ref}.xlsx"'
    )

    wb.save(response)
    return response